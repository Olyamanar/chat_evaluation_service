import io
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from models import ChatEvaluation, Chat, CRITERIA


HEADER_FILL = PatternFill(start_color="6C5CE7", end_color="6C5CE7", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="6C5CE7")
SUBTITLE_FONT = Font(name="Arial", size=11, bold=True, color="6C5CE7")
NORMAL_FONT = Font(name="Arial", size=10)
WRAP_ALIGN = Alignment(wrap_text=True, vertical="top")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="BDC3C7"),
    right=Side(style="thin", color="BDC3C7"),
    top=Side(style="thin", color="BDC3C7"),
    bottom=Side(style="thin", color="BDC3C7"),
)

SCORE_COLORS = {
    1: PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid"),
    3: PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid"),
    5: PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid"),
}

TOTAL_COLORS = {
    0: PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid"),
    30: PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid"),
    100: PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid"),
}


def export_to_excel(
    evaluations: List[ChatEvaluation],
    chats: List[Chat],
    employee_name: str
) -> bytes:
    wb = Workbook()

    chat_map = {c.id: c for c in chats}

    _write_summary_sheet(wb, evaluations, employee_name)
    _write_details_sheet(wb, evaluations, chat_map)
    _write_recommendations_sheet(wb, evaluations)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _write_summary_sheet(wb: Workbook, evaluations: List[ChatEvaluation], employee_name: str):
    ws = wb.active
    ws.title = "Сводка"

    ws.merge_cells("A1:F1")
    ws["A1"] = f"Оценка софт-скиллов: {employee_name}"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    headers = [
        "№", "ID диалога", "Дата",
        "Софт-скиллы (1-5)",
        "Итого баллов", "Резюме"
    ]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER

    for idx, ev in enumerate(evaluations, 1):
        row = idx + 3
        ws.cell(row=row, column=1, value=idx).alignment = CENTER_ALIGN
        ws.cell(row=row, column=2, value=ev.chat_id).alignment = CENTER_ALIGN

        date_val = ev.date or ""
        ws.cell(row=row, column=3, value=date_val).alignment = CENTER_ALIGN

        for ci, cs in enumerate(ev.criteria_scores):
            cell = ws.cell(row=row, column=4 + ci, value=cs.score)
            cell.alignment = CENTER_ALIGN
            cell.fill = SCORE_COLORS.get(cs.score, PatternFill())
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.border = THIN_BORDER

        total_cell = ws.cell(row=row, column=5, value=ev.total_score)
        total_cell.alignment = CENTER_ALIGN
        total_cell.fill = TOTAL_COLORS.get(ev.total_score, PatternFill())
        total_cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        total_cell.border = THIN_BORDER

        ws.cell(row=row, column=6, value=ev.summary).alignment = WRAP_ALIGN

    last_row = len(evaluations) + 4
    ws.cell(row=last_row, column=1, value="Средний балл:").font = SUBTITLE_FONT
    avg = sum(e.total_score for e in evaluations) / len(evaluations) if evaluations else 0
    avg_cell = ws.cell(row=last_row, column=2, value=round(avg, 1))
    avg_cell.font = Font(name="Arial", size=12, bold=True)

    col_widths = [5, 12, 12, 16, 14, 50]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_details_sheet(wb: Workbook, evaluations: List[ChatEvaluation], chat_map: dict):
    ws = wb.create_sheet("Подробно")

    ws.merge_cells("A1:E1")
    ws["A1"] = "Подробная оценка по каждому диалогу"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30

    current_row = 3
    for idx, ev in enumerate(evaluations, 1):
        ws.merge_cells(f"A{current_row}:E{current_row}")
        ws.cell(row=current_row, column=1, value=f"Диалог #{ev.chat_id} — {ev.employee_name} — Итого: {ev.total_score}/100")
        ws.cell(row=current_row, column=1).font = SUBTITLE_FONT
        ws.cell(row=current_row, column=1).fill = PatternFill(start_color="ECF0F1", end_color="ECF0F1", fill_type="solid")
        current_row += 1

        chat = chat_map.get(ev.chat_id)
        if chat:
            ws.merge_cells(f"A{current_row}:E{current_row}")
            ws.cell(row=current_row, column=1, value="Текст диалога:")
            ws.cell(row=current_row, column=1).font = Font(name="Arial", size=10, bold=True, italic=True)
            current_row += 1
            for msg in chat.messages:
                if msg.role == "bot":
                    continue
                role_label = "Сотрудник" if msg.role == "employee" else "Клиент"
                ws.merge_cells(f"A{current_row}:E{current_row}")
                ws.cell(row=current_row, column=1, value=f"  {role_label}: {msg.text}")
                ws.cell(row=current_row, column=1).font = NORMAL_FONT
                ws.cell(row=current_row, column=1).alignment = WRAP_ALIGN
                ws.row_dimensions[current_row].height = max(15, len(msg.text) // 3)
                current_row += 1
            current_row += 1

        headers = ["Критерий", "Оценка", "Обоснование"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = CENTER_ALIGN
            cell.border = THIN_BORDER
        current_row += 1

        for cs in ev.criteria_scores:
            ws.cell(row=current_row, column=1, value=cs.name).font = NORMAL_FONT
            score_cell = ws.cell(row=current_row, column=2, value=cs.score)
            score_cell.alignment = CENTER_ALIGN
            score_cell.fill = SCORE_COLORS.get(cs.score, PatternFill())
            score_cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            score_cell.border = THIN_BORDER
            ws.cell(row=current_row, column=3, value=cs.justification).font = NORMAL_FONT
            ws.cell(row=current_row, column=3).alignment = WRAP_ALIGN
            ws.row_dimensions[current_row].height = max(30, len(cs.justification) // 4)
            current_row += 1

        current_row += 2

    col_widths = [60, 10, 80, 30, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_recommendations_sheet(wb: Workbook, evaluations: List[ChatEvaluation]):
    ws = wb.create_sheet("Рекомендации")

    ws.merge_cells("A1:D1")
    ws["A1"] = "Рекомендации для улучшения работы"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30

    all_strengths = []
    all_weaknesses = []
    all_recommendations = []

    for ev in evaluations:
        all_strengths.extend(ev.strengths)
        all_weaknesses.extend(ev.weaknesses)
        all_recommendations.extend(ev.recommendations)

    current_row = 3

    ws.merge_cells(f"A{current_row}:D{current_row}")
    ws.cell(row=current_row, column=1, value="Сильные стороны")
    ws.cell(row=current_row, column=1).font = Font(name="Arial", size=12, bold=True, color="27AE60")
    current_row += 1

    unique_strengths = list(dict.fromkeys(all_strengths))
    if unique_strengths:
        for s in unique_strengths[:20]:
            ws.cell(row=current_row, column=1, value=f"  • {s}").font = NORMAL_FONT
            ws.cell(row=current_row, column=1).alignment = WRAP_ALIGN
            current_row += 1
    else:
        ws.cell(row=current_row, column=1, value="  Не выявлены").font = Font(name="Arial", size=10, italic=True)
        current_row += 1

    current_row += 1

    ws.merge_cells(f"A{current_row}:D{current_row}")
    ws.cell(row=current_row, column=1, value="Зоны роста (слабые стороны)")
    ws.cell(row=current_row, column=1).font = Font(name="Arial", size=12, bold=True, color="E74C3C")
    current_row += 1

    unique_weaknesses = list(dict.fromkeys(all_weaknesses))
    if unique_weaknesses:
        for w in unique_weaknesses[:20]:
            ws.cell(row=current_row, column=1, value=f"  • {w}").font = NORMAL_FONT
            ws.cell(row=current_row, column=1).alignment = WRAP_ALIGN
            current_row += 1
    else:
        ws.cell(row=current_row, column=1, value="  Не выявлены").font = Font(name="Arial", size=10, italic=True)
        current_row += 1

    current_row += 1

    ws.merge_cells(f"A{current_row}:D{current_row}")
    ws.cell(row=current_row, column=1, value="Рекомендации для сотрудника")
    ws.cell(row=current_row, column=1).font = Font(name="Arial", size=12, bold=True, color="2980B9")
    current_row += 1

    unique_recs = list(dict.fromkeys(all_recommendations))
    if unique_recs:
        for r in unique_recs[:20]:
            ws.cell(row=current_row, column=1, value=f"  → {r}").font = NORMAL_FONT
            ws.cell(row=current_row, column=1).alignment = WRAP_ALIGN
            current_row += 1
    else:
        ws.cell(row=current_row, column=1, value="  Нет рекомендаций").font = Font(name="Arial", size=10, italic=True)
        current_row += 1

    ws.column_dimensions["A"].width = 80
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 30
